"""Exporta un listado tabular a CSV, XLSX o PDF (feature 013).

Genérico: recibe columnas `[(clave, título), ...]` y filas `list[dict]`, devuelve
`(bytes, media_type, extensión)`. Lo usan el Catálogo y Riesgo de Fuga.
"""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from decimal import Decimal
from typing import Literal

Formato = Literal["csv", "xlsx", "pdf"]

_MEDIA = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


def _val(v) -> str:
    if v is None:
        return ""
    if isinstance(v, Decimal | float):
        return f"{v:.2f}".rstrip("0").rstrip(".")
    if isinstance(v, date | datetime):
        return v.isoformat()
    return str(v)


def exportar(
    formato: Formato,
    *,
    titulo: str,
    columnas: list[tuple[str, str]],
    filas: list[dict],
) -> tuple[bytes, str, str]:
    if formato == "csv":
        return _csv(columnas, filas), _MEDIA["csv"], "csv"
    if formato == "xlsx":
        return _xlsx(titulo, columnas, filas), _MEDIA["xlsx"], "xlsx"
    if formato == "pdf":
        return _pdf(titulo, columnas, filas), _MEDIA["pdf"], "pdf"
    raise ValueError(f"formato no soportado: {formato}")


def _csv(columnas: list[tuple[str, str]], filas: list[dict]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([t for _, t in columnas])
    for f in filas:
        w.writerow([_val(f.get(k)) for k, _ in columnas])
    return buf.getvalue().encode("utf-8-sig")


def _xlsx(titulo: str, columnas: list[tuple[str, str]], filas: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = titulo[:31]
    cab = Font(bold=True, color="FFFFFF")
    relleno = PatternFill("solid", fgColor="0A3632")
    for col, (_, t) in enumerate(columnas, 1):
        c = ws.cell(row=1, column=col, value=t)
        c.font = cab
        c.fill = relleno
    for r, fila in enumerate(filas, 2):
        for col, (k, _) in enumerate(columnas, 1):
            v = fila.get(k)
            ws.cell(row=r, column=col, value=float(v) if isinstance(v, Decimal) else v)
    for col, (k, t) in enumerate(columnas, 1):
        ancho = max(len(t), *(len(_val(f.get(k))) for f in filas)) if filas else len(t)
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = min(48, ancho + 3)
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _pdf(titulo: str, columnas: list[tuple[str, str]], filas: list[dict]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=landscape(A4),
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=14 * mm,
        title=titulo,
    )
    estilos = getSampleStyleSheet()
    celda = estilos["BodyText"]
    celda.fontSize = 7
    celda.leading = 9
    data = [[t for _, t in columnas]]
    for f in filas:
        data.append([Paragraph(_val(f.get(k)), celda) for k, _ in columnas])
    tabla = Table(data, repeatRows=1)
    tabla.setStyle(
        TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0A3632")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTSIZE", (0, 0), (-1, 0), 7),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F4F6")]),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#C0C8C6")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ])
    )
    doc.build([
        Paragraph(f"<b>{titulo}</b>", estilos["Title"]),
        Paragraph(
            f"Marzú Retail Group · {datetime.now():%d-%m-%Y %H:%M} · {len(filas)} filas",
            estilos["Normal"],
        ),
        Spacer(1, 6 * mm),
        tabla,
    ])
    return out.getvalue()
